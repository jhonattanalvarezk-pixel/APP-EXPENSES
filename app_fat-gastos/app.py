from datetime import datetime
import json
import os
import re
import tempfile
import time
from google import genai
from google.genai import types
import openpyxl
import pandas as pd
import io
import streamlit as st
import base64
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaFileUpload, MediaIoBaseDownload
from PIL import Image, ImageEnhance
import streamlit.components.v1 as components

def get_drive_service():
    """Autentica con Google Drive usando las credenciales OAuth y el Refresh Token."""
    creds = Credentials(
        None,
        refresh_token=st.secrets["OAUTH_REFRESH_TOKEN"],
        client_id=st.secrets["OAUTH_CLIENT_ID"],
        client_secret=st.secrets["OAUTH_CLIENT_SECRET"],
        token_uri="https://oauth2.googleapis.com/token",
    )
    
    if not creds.valid:
        creds.refresh(Request())
        
    return build("drive", "v3", credentials=creds)

def subir_archivo_a_drive(file_path, file_name):
    """Sube un archivo local a la carpeta de Google Drive."""
    try:
        service = get_drive_service()
        folder_id = st.secrets.get('DRIVE_FOLDER_ID')
        
        query = f"name = '{file_name}' and '{folder_id}' in parents and trashed = false"
        results = service.files().list(q=query, fields="files(id)").execute()
        files = results.get("files", [])
        
        media = MediaFileUpload(file_path, resumable=True)
        if files:
            file_id = files[0]['id']
            file = service.files().update(
                fileId=file_id,
                media_body=media
            ).execute()
        else:
            file_metadata = {
                'name': file_name,
                'parents': [folder_id]
            }
            file = service.files().create(
                body=file_metadata,
                media_body=media,
                fields='id'
            ).execute()
        st.success(f"¡Subido con éxito a Drive: {file_name}!")
        return file.get('id')
    except Exception as e:
        st.error(f"❌ ERROR TÉCNICO DE GOOGLE DRIVE: {e}")
        return None

def descargar_archivo_de_drive(nombre_archivo, destino_path):
    """Busca el archivo por su nombre en Google Drive y lo descarga a local si no existe."""
    try:
        service = get_drive_service()
        folder_id = st.secrets.get('DRIVE_FOLDER_ID')
        
        query = f"name = '{nombre_archivo}' and '{folder_id}' in parents and trashed = false"
        results = service.files().list(q=query, fields="files(id, name)").execute()
        files = results.get("files", [])
        
        if not files:
            return False
            
        file_id = files[0]['id']
        request = service.files().get_media(fileId=file_id)
        fh = io.BytesIO()
        downloader = MediaIoBaseDownload(fh, request)
        done = False
        while not done:
            status, done = downloader.next_chunk()
            
        fh.seek(0)
        with open(destino_path, "wb") as f:
            f.write(fh.read())
        return True
    except Exception as e:
        print(f"Error descargando desde Google Drive: {e}")
        return False

def mostrar_archivo_inteligente(file_path, archivo_nombre):
    """Muestra la imagen o PDF convertido directamente en pantalla con un contenedor limpio."""
    if not os.path.exists(file_path):
        with st.spinner("Recuperando documento..."):
            encontrado = descargar_archivo_de_drive(archivo_nombre, file_path)
            if not encontrado:
                st.error("El archivo no se encuentra disponible localmente.")
                return

    if archivo_nombre.lower().endswith(".pdf"):
        try:
            with open(file_path, "rb") as f:
                pdf_bytes = f.read()
            
            st.download_button(
                label="📥 Descargar / Abrir PDF original",
                data=pdf_bytes,
                file_name=archivo_nombre,
                mime="application/pdf",
                key=f"dl_visor_{archivo_nombre}",
                type="primary"
            )
            
            import fitz
            doc = fitz.open(file_path)
            page = doc[0]
            pix = page.get_pixmap(dpi=150)
            img_path = file_path + "_vis.png"
            pix.save(img_path)
            st.image(img_path, caption="Vista previa del documento guardado", use_container_width=True)
            
        except Exception as e:
            st.error(f"No se pudo renderizar la vista previa del PDF: {e}")
    else:
        # Contenedor limpio para imágenes largas
        st.image(file_path, caption="Documento / Ticket", use_container_width=True)

def aclarar_y_optimizar_imagen(image_bytes):
    """Comprime, redimensiona y aplica alto contraste automáticamente."""
    try:
        img = Image.open(io.BytesIO(image_bytes))
        
        if img.mode in ("RGBA", "P"):
            img = img.convert("RGB")
            
        max_width = 1200
        if img.width > max_width:
            w_percent = (max_width / float(img.width))
            h_size = int(float(img.height) * float(w_percent))
            img = img.resize((max_width, h_size), Image.Resampling.LANCZOS)
            
        enhancer_contrast = ImageEnhance.Contrast(img)
        img = enhancer_contrast.enhance(1.5)
        
        enhancer_sharpness = ImageEnhance.Sharpness(img)
        img = enhancer_sharpness.enhance(1.6)
        
        output_io = io.BytesIO()
        img.save(output_io, format="JPEG", quality=85, optimize=True)
        return output_io.getvalue()
    except Exception:
        return image_bytes

# Configuración de la página
st.set_page_config(
    page_title="Gestión de Facturas y Gastos Pro", page_icon="🧾", layout="wide"
)

st.markdown(
    """
    <style>
    .stApp {
        background-color: #0F172A;
        color: #F8FAFC;
    }
    [data-testid="stSidebar"] {
        background-color: #1E293B !important;
        border-right: 1px solid #334155;
    }
    [data-testid="stMetricValue"] {
        font-size: 1.8rem !important;
        font-weight: 700 !important;
        color: #38BDF8 !important;
    }
    div[data-testid="metric-container"] {
        background-color: #1E293B;
        border: 1px solid #334155;
        padding: 15px;
        border-radius: 12px;
        box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.1);
    }
    div.stButton > button[kind="primary"] {
        background: linear-gradient(135deg, #2563EB 0%, #1D4ED8 100%);
        color: white;
        border-radius: 8px;
        font-weight: 600;
        border: none;
        padding: 0.6rem 1.2rem;
    }
    .stForm {
        background-color: #1E293B;
        border: 1px solid #334155 !important;
        padding: 20px;
        border-radius: 12px;
    }
    /* Estilo para contener imágenes largas de tickets de forma ordenada */
    [data-testid="stImage"] img {
        border-radius: 8px;
        max-height: 500px;
        object-fit: contain;
        background-color: #0B0F19;
    }
    </style>
""",
    unsafe_allow_html=True,
)

LOGO_PATH = "logo.png"

if os.path.exists(LOGO_PATH):
  st.sidebar.image(LOGO_PATH, use_container_width=True)
else:
  st.sidebar.info(f"💡 Coloca tu archivo '{LOGO_PATH}' en la carpeta del app.")

st.sidebar.title("Configuración")
api_key = st.sidebar.text_input(
    "Gemini API Key", type="password", value=os.environ.get("GEMINI_API_KEY", "")
)

EXCEL_FILE = "registro_facturas.xlsx"
PDF_DIR = "archivos_facturas"

os.makedirs(PDF_DIR, exist_ok=True)

def inicializar_excel():
  if not os.path.exists(EXCEL_FILE):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Facturas"
    headers = [
        "ID",
        "Fecha",
        "Proveedor",
        "Origen",
        "Categoria",
        "Base Imponible",
        "IVA (%)",
        "Cuota IVA",
        "Total",
        "Archivo Original",
    ]
    ws.append(headers)
    wb.save(EXCEL_FILE)

inicializar_excel()

def limpiar_json_string(texto: str) -> str:
  texto = re.sub(r"^```json\s*", "", texto, flags=re.MULTILINE)
  texto = re.sub(r"^```\s*", "", texto, flags=re.MULTILINE)
  texto = re.sub(r"```$", "", texto, flags=re.MULTILINE)
  return texto.strip()

def extraer_datos_factura(file_bytes, mime_type, api_key):
  client = genai.Client(api_key=api_key)

  prompt = """
    Analiza este documento de factura o ticket y extrae los siguientes datos en formato JSON estricto:
    - fecha: Fecha de la factura en formato AAAA-MM-DD (si no se encuentra, pon la fecha actual).
    - proveedor: Nombre de la tienda, empresa o emisor de la factura.
    - categoria: Clasifica el gasto en una de estas opciones: Materiales, Transporte, Herramientas, Varios.
    - base_imponible: Numero decimal con la base imponible total.
    - iva_porcentaje: Numero entero o decimal con el tipo de IVA principal (ej: 21, 10, 4).
    - cuota_iva: Numero decimal con el importe total del IVA.
    - total: Numero decimal con el importe total de la factura.
    
    Devuelve UNICAMENTE un JSON valido con estas claves: fecha, proveedor, categoria, base_imponible, iva_porcentaje, cuota_iva, total.
    """

  ext = ".pdf" if "pdf" in mime_type else ".jpg"

  with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp_file:
    tmp_file.write(file_bytes)
    tmp_path = tmp_file.name

  uploaded_file = None
  try:
    uploaded_file = client.files.upload(file=tmp_path)

    if "pdf" in mime_type:
      while uploaded_file.state.name == "PROCESSING":
        time.sleep(1)
        uploaded_file = client.files.get(name=uploaded_file.name)

      if uploaded_file.state.name == "FAILED":
        raise Exception("Google Gemini no pudo procesar este archivo PDF.")

    modelos = ["gemini-2.5-flash", "gemini-2.0-flash", "gemini-3.8-flash"]
    ultimo_error = None

    for modelo in modelos:
      for intento in range(4):
        try:
          response = client.models.generate_content(
              model=modelo,
              contents=[uploaded_file, prompt],
              config=types.GenerateContentConfig(
                  response_mime_type="application/json"
              ),
          )
          return response.text
        except Exception as e:
          ultimo_error = e
          msg_error = str(e)

          if "404" in msg_error or "NOT_FOUND" in msg_error:
            break

          if (
              "503" in msg_error
              or "UNAVAILABLE" in msg_error
              or "RESOURCE_EXHAUSTED" in msg_error
              or "429" in msg_error
          ):
            time.sleep(2**intento)
            continue
          else:
            break

    if ultimo_error:
      raise ultimo_error

  finally:
    if os.path.exists(tmp_path):
      os.remove(tmp_path)
    if uploaded_file:
      try:
        client.files.delete(name=uploaded_file.name)
      except Exception:
        pass

  raise Exception(
      "Los servidores de Gemini estan con alta demanda. Reintentalo en unos instantes."
  )

# --- CABECERA PRINCIPAL ---
col_head1, col_head2 = st.columns([0.8, 0.2])
with col_head1:
  st.title("Gestor de Facturas y Gastos Pro")
with col_head2:
  if os.path.exists(LOGO_PATH):
    st.image(LOGO_PATH, width=120)

menu = st.sidebar.selectbox(
    "Menú",
    [
        "Nueva Factura (Individual/Lote)",
        "Listado y Edición Interactiva",
        "Resumen Trimestral y Exportación",
    ],
)

# 1. OPCION: NUEVA FACTURA
if menu == "Nueva Factura (Individual/Lote)":
  st.header("Carga de Facturas y Tickets")

  st.markdown("📂 **Centro de Captura:** Sube o toma fotos de tus tickets y facturas (incluso los más largos). El sistema los comprimirá y organizará automáticamente.")

  uploaded_files = st.file_uploader(
      "Toca aquí para Tomar Foto, abrir la Cámara, seleccionar de la Galería o subir PDFs",
      type=["png", "jpg", "jpeg", "webp", "pdf"],
      accept_multiple_files=True,
      key="selector_universal_facturas"
  )

  uploaded_files_list = []

  if uploaded_files:
    for uf in uploaded_files:
      if uf.type.startswith("image/"):
        bytes_mejorados = aclarar_y_optimizar_imagen(uf.getvalue())
        class ProcessedImage:
          def __init__(self, data, name, mime):
            self.data = data
            self.name = name
            self.type = mime
          def getvalue(self):
            return self.data
        uploaded_files_list.append(ProcessedImage(bytes_mejorados, uf.name, "image/jpeg"))
      else:
        uploaded_files_list.append(uf)

  if uploaded_files_list and not api_key:
    st.warning("Por favor, introduce tu Gemini API Key en la barra lateral.")

  elif uploaded_files_list and api_key:
    if st.button("Procesar Facturas con IA", type="primary"):
      for uploaded_file in uploaded_files_list:
        with st.spinner(f"Procesando {uploaded_file.name}..."):
          file_bytes = uploaded_file.getvalue()
          mime_type = uploaded_file.type

          timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
          safe_filename = f"{timestamp}_{uploaded_file.name}"
          file_path = os.path.join(PDF_DIR, safe_filename)
          
          with open(file_path, "wb") as f:
            f.write(file_bytes)

          try:
            json_str = extraer_datos_factura(file_bytes, mime_type, api_key)
            json_str_limpio = limpiar_json_string(json_str)
            datos = json.loads(json_str_limpio)

            st.session_state[f"datos_{safe_filename}"] = {
                "datos": datos,
                "safe_filename": safe_filename,
                "file_path": file_path,
                "mime_type": mime_type,
                "file_bytes": file_bytes,
            }
            st.success(f"¡Procesado con éxito: {uploaded_file.name}!")

          except Exception as e:
            st.error(f"Error procesando {uploaded_file.name}: {e}")

  keys_to_show = [
      k for k in st.session_state.keys() if k.startswith("datos_")
  ]

  if keys_to_show:
    st.subheader("Facturas pendientes de guardar en Excel")

    if len(keys_to_show) > 1:
      if st.button(
          f"Guardar TODAS ({len(keys_to_show)}) en Excel de golpe",
          type="secondary",
      ):
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active

        for key in list(keys_to_show):
          item = st.session_state[key]
          datos = item["datos"]
          safe_filename = item["safe_filename"]
          file_path = item["file_path"]
          file_bytes = item.get("file_bytes")

          if not os.path.exists(file_path) and file_bytes:
            with open(file_path, "wb") as f:
              f.write(file_bytes)

          nuevo_id = ws.max_row
          ws.append([
              nuevo_id,
              str(datos.get("fecha", datetime.now().strftime("%Y-%m-%d"))),
              str(datos.get("proveedor", "Desconocido")),
              "Cámara / Archivo Móvil",
              str(datos.get("categoria", "Materiales")),
              float(datos.get("base_imponible") or 0.00),
              float(datos.get("iva_porcentaje") or 21.0),
              float(datos.get("cuota_iva") or 0.00),
              float(datos.get("total") or 0.00),
              safe_filename,
          ])
          
          subir_archivo_a_drive(file_path, safe_filename)
          del st.session_state[key]

        wb.save(EXCEL_FILE)
        subir_archivo_a_drive(EXCEL_FILE, EXCEL_FILE)
        st.success("¡Todas las facturas del lote han sido guardadas en Excel!")
        time.sleep(1)
        st.rerun()

    for key in keys_to_show:
      item = st.session_state[key]
      datos = item["datos"]
      
      if isinstance(datos, list):
          datos = datos[0] if len(datos) > 0 else {}
          
      safe_filename = item["safe_filename"]
      file_path = item["file_path"]
      mime_type = item["mime_type"]
      file_bytes = item.get("file_bytes")

      if not os.path.exists(file_path) and file_bytes is not None:
        with open(file_path, "wb") as f:
          f.write(file_bytes)

      fecha_str = datos.get("fecha", "")
      try:
        fecha_valida = datetime.strptime(fecha_str, "%Y-%m-%d").date()
      except (ValueError, TypeError):
        fecha_valida = datetime.now().date()

      with st.expander(f"Revisar: {safe_filename.split('_', 2)[-1]}", expanded=True):
        col_form, col_visor = st.columns([1.1, 0.9])

        with col_visor:
          st.markdown("**📸 Previsualización del Documento**")
          if os.path.exists(file_path):
            if safe_filename.lower().endswith(".pdf"):
              st.info("📄 Documento PDF convertido para previsualización:")
              with open(file_path, "rb") as f:
                pdf_bytes = f.read()
              st.download_button(
                  label="📥 Descargar PDF original",
                  data=pdf_bytes,
                  file_name=safe_filename,
                  mime="application/pdf",
                  key=f"dl_local_{safe_filename}",
                  type="primary"
              )
              try:
                import fitz
                doc = fitz.open(file_path)
                page = doc[0]
                pix = page.get_pixmap(dpi=150)
                img_path = file_path + ".png"
                pix.save(img_path)
                st.image(img_path, caption="Vista previa de la factura (PDF)", use_container_width=True)
              except Exception:
                st.warning("No se pudo generar la vista previa automática del PDF.")
            else:
              st.image(file_path, caption="Ticket optimizado para lectura", use_container_width=True)
          else:
            st.warning("El archivo local no está disponible para previsualizar.")

        with col_form:
          with st.form(key=f"form_{safe_filename}"):
            f_fecha = st.date_input("Fecha", value=fecha_valida)
            f_proveedor = st.text_input(
                "Proveedor", value=str(datos.get("proveedor", ""))
            )
            cat_detectada = str(datos.get("categoria", "Materiales"))
            categorias_validas = [
                "Materiales",
                "Transporte",
                "Herramientas",
                "Varios",
            ]
            index_cat = (
                categorias_validas.index(cat_detectada)
                if cat_detectada in categorias_validas
                else 0
            )
            f_categoria = st.selectbox(
                "Categoría", categorias_validas, index=index_cat
            )

            f_base = st.number_input(
                "Base Imponible (EUR)",
                value=float(datos.get("base_imponible") or 0.00),
                format="%.2f",
            )
            f_iva_p = st.number_input(
                "IVA (%)",
                value=float(datos.get("iva_porcentaje") or 21.0),
                format="%.1f",
            )
            f_cuota = st.number_input(
                "Cuota IVA (EUR)",
                value=float(datos.get("cuota_iva") or 0.00),
                format="%.2f",
            )
            f_total = st.number_input(
                "Total (EUR)",
                value=float(datos.get("total") or 0.00),
                format="%.2f",
            )

            submitted = st.form_submit_button("Guardar Factura en Excel")
            if submitted:
              if not os.path.exists(file_path) and file_bytes:
                with open(file_path, "wb") as f:
                  f.write(file_bytes)

              wb = openpyxl.load_workbook(EXCEL_FILE)
              ws = wb.active
              nuevo_id = ws.max_row
              ws.append([
                  nuevo_id,
                  str(f_fecha),
                  f_proveedor,
                  "Cámara / Archivo Móvil",
                  f_categoria,
                  f_base,
                  f_iva_p,
                  f_cuota,
                  f_total,
                  safe_filename,
              ])
              wb.save(EXCEL_FILE)

              subir_archivo_a_drive(file_path, safe_filename)
              subir_archivo_a_drive(EXCEL_FILE, EXCEL_FILE)

              del st.session_state[key]
              st.success(
                  f"Factura de {f_proveedor} guardada en Excel con éxito."
              )
              time.sleep(1)
              st.rerun()

# 2. OPCION: LISTADO Y EDICION INTERACTIVA
elif menu == "Listado y Edición Interactiva":
  st.header("Listado, Edición y Registro de Facturas")
  if os.path.exists(EXCEL_FILE):
    df = pd.read_excel(EXCEL_FILE)
    if not df.empty:
      st.markdown(
          "Puedes **editar valores directamente en la tabla** o seleccionar las"
          " casillas para guardar o eliminar filas."
      )

      edited_df = st.data_editor(
          df,
          num_rows="dynamic",
          use_container_width=True,
          key="tabla_interactiva",
      )

      if st.button("Guardar Cambios Realizados en la Tabla"):
        edited_df.to_excel(EXCEL_FILE, index=False)
        subir_archivo_a_drive(EXCEL_FILE, EXCEL_FILE)
        st.success("¡El archivo Excel ha sido actualizado correctamente!")
        time.sleep(1)
        st.rerun()

      st.divider()
      st.subheader("Visor de Archivo Guardado")
      archivos_disponibles = [
          f for f in df["Archivo Original"].dropna().unique()
      ]
      if archivos_disponibles:
        archivo_sel = st.selectbox(
            "Selecciona una factura para consultar el documento:",
            archivos_disponibles,
        )
        ruta_archivo = os.path.join(PDF_DIR, archivo_sel)
        mostrar_archivo_inteligente(ruta_archivo, archivo_sel)
    else:
      st.info("No hay facturas registradas todavía.")
  else:
    st.info("El archivo Excel aún no ha sido creado.")

# 3. OPCION: RESUMEN TRIMESTRAL Y EXPORTACION COMPLETA
elif menu == "Resumen Trimestral y Exportación":
  st.header("Resumen Trimestral, IVA Soportado y Exportación")
  if os.path.exists(EXCEL_FILE):
    df = pd.read_excel(EXCEL_FILE)
    if not df.empty:
      total_base = df["Base Imponible"].sum()
      total_iva = df["Cuota IVA"].sum()
      total_general = df["Total"].sum()

      col1, col2, col3 = st.columns(3)
      col1.metric("Base Imponible Total", f"{total_base:.2f} EUR")
      col2.metric("IVA Soportado Total", f"{total_iva:.2f} EUR")
      col3.metric("Gasto Total", f"{total_general:.2f} EUR")

      st.subheader("Desglose por Categoría")
      col_cat = "Categoria" if "Categoria" in df.columns else "Categoría"
      resumen_cat = (
          df.groupby(col_cat)[["Base Imponible", "Cuota IVA", "Total"]]
          .sum()
          .reset_index()
      )
      st.dataframe(resumen_cat, use_container_width=True)

      st.divider()
      st.subheader("Exportación de Informes")

      with open(EXCEL_FILE, "rb") as f:
        excel_bytes = f.read()

      st.download_button(
          label="Descargar Excel de Registro de Facturas (.xlsx)",
          data=excel_bytes,
          file_name=f"registro_facturas_{datetime.now().strftime('%Y%m%d')}.xlsx",
          mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
          type="primary",
      )
    else:
      st.info("Aún no hay datos para resumir.")
  else:
    st.info("No hay registros disponibles.")